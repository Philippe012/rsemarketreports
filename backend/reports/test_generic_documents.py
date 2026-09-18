import os

import openpyxl
from django.conf import settings
from django.test import SimpleTestCase, TestCase

from reports.models import Report
from reports.tests import AuthenticatedTestCase
from services.documents.classification import classify_document
from services.documents.schema_inference import infer_column

SAMPLE_DIR = os.path.join(settings.BASE_DIR.parent, 'sample_data')
SALES_CSV = os.path.join(SAMPLE_DIR, 'generic_sales_sample.csv')
HR_XLSX = os.path.join(SAMPLE_DIR, 'generic_hr_sample.xlsx')
OPERATIONS_DOCX = os.path.join(SAMPLE_DIR, 'generic_operations_sample.docx')
NOTES_TXT = os.path.join(SAMPLE_DIR, 'generic_notes_sample.txt')
PROSE_ONLY_TXT = os.path.join(SAMPLE_DIR, 'generic_prose_only_sample.txt')
OPERATIONS_PDF = os.path.join(SAMPLE_DIR, 'generic_operations_sample.pdf')
ORDERS_JSON = os.path.join(SAMPLE_DIR, 'generic_orders_sample.json')


class SchemaInferenceTests(SimpleTestCase):
    def test_infers_currency_from_name_and_values(self):
        column = infer_column('Revenue', [4800000, 3400000, 5600000])
        self.assertEqual(column['semantic_type'], 'currency')

    def test_infers_percentage_from_percent_sign(self):
        column = infer_column('Discount', ['5%', '0%', '10%'])
        self.assertEqual(column['semantic_type'], 'percentage')

    def test_infers_date_column(self):
        column = infer_column('Date', ['2025-01-15', '2025-02-10', '2025-03-05'])
        self.assertEqual(column['semantic_type'], 'date')

    def test_infers_category_for_low_cardinality_text(self):
        column = infer_column('Region', ['East Africa', 'West Africa', 'East Africa', 'Central Africa'] * 3)
        self.assertEqual(column['semantic_type'], 'category')

    def test_infers_identifier_for_high_cardinality_coded_column(self):
        column = infer_column('Employee ID', [f'E{i:03d}' for i in range(20)])
        self.assertEqual(column['semantic_type'], 'identifier')

    def test_never_invents_a_type_for_an_empty_column(self):
        column = infer_column('Notes', [])
        self.assertEqual(column['non_null_count'], 0)
        self.assertEqual(column['confidence'], 'low')

    def test_tracks_unparsed_values_separately_from_missing_ones(self):
        # "N/A" is present but not a number — distinct from a value that was
        # never provided at all, so it should count as unparsed, not missing.
        # (numeric_ratio must stay >= 0.8 for the column to be read as numeric at all.)
        column = infer_column('Revenue', [4800000, 'N/A', 3400000, 5600000, 2100000, 1900000], total_count=7)
        self.assertEqual(column['stats']['unparsed_count'], 1)
        self.assertEqual(column['stats']['missing_count'], 2)  # 1 unparsed + 1 truly absent row

    def test_tracks_negative_values(self):
        column = infer_column('Revenue', [4800000, -200000, 3400000])
        self.assertEqual(column['stats']['negative_count'], 1)

    def test_identifier_stats_expose_uniqueness_for_duplicate_detection(self):
        # Duplicate detection only makes sense once a column is "almost
        # entirely" unique (>90%) — that's the same bar infer_column uses to
        # call it an identifier in the first place.
        ids = [f'E{i:03d}' for i in range(1, 11)] + ['E001']
        column = infer_column('Employee ID', ids)
        self.assertEqual(column['semantic_type'], 'identifier')
        self.assertEqual(column['stats']['unique_count'], 10)


class GenericValidationWarningsTests(SimpleTestCase):
    def _dataset(self, columns, rows):
        return {
            'name': 'Data', 'source': 'test', 'row_count': len(rows),
            'duplicate_row_count': 0, 'columns': columns, 'rows': rows,
        }

    def test_warns_about_unparsed_numeric_values(self):
        from services.documents.validate import validate_document

        column = infer_column('Revenue', [4800000, 'N/A', 3400000, 5600000, 2100000, 1900000])
        document = {'datasets': [self._dataset([column], [{}] * 6)], 'sections': [], 'figures': []}
        _, warnings = validate_document(document)
        self.assertTrue(any('did not look like a valid number' in w and 'Revenue' in w for w in warnings))

    def test_warns_about_negative_currency_values(self):
        from services.documents.validate import validate_document

        column = infer_column('Revenue', [4800000, -200000, 3400000])
        document = {'datasets': [self._dataset([column], [{}, {}, {}])], 'sections': [], 'figures': []}
        _, warnings = validate_document(document)
        self.assertTrue(any('negative' in w and 'Revenue' in w for w in warnings))

    def test_warns_about_duplicate_identifier_values(self):
        from services.documents.validate import validate_document

        ids = [f'E{i:03d}' for i in range(1, 11)] + ['E001']
        column = infer_column('Employee ID', ids)
        document = {'datasets': [self._dataset([column], [{}] * 11)], 'sections': [], 'figures': []}
        _, warnings = validate_document(document)
        self.assertTrue(any('duplicate' in w.lower() and 'Employee ID' in w for w in warnings))

    def test_clean_data_raises_no_new_warnings(self):
        from services.documents.validate import validate_document

        column = infer_column('Revenue', [4800000, 3400000, 5600000])
        document = {'datasets': [self._dataset([column], [{}, {}, {}])], 'sections': [], 'figures': []}
        _, warnings = validate_document(document)
        self.assertEqual(warnings, [])


class EntityExtractionTests(SimpleTestCase):
    def test_extracts_currency_and_percent_mentions(self):
        from services.documents.entities import extract_entities

        text = "Revenue grew 12% to reach $4.5 million, with costs up 3%."
        entities = extract_entities(text)
        self.assertIn('$4.5 million', entities['numbers'])
        self.assertIn('12%', entities['numbers'])

    def test_extracts_dates(self):
        from services.documents.entities import extract_entities

        entities = extract_entities("The program launched in March 2024 and ended on 12/05/2025.")
        self.assertIn('March 2024', entities['dates'])
        self.assertIn('12/05/2025', entities['dates'])

    def test_keywords_require_repetition(self):
        from services.documents.entities import extract_entities

        entities = extract_entities("The clinic served patients. The clinic was busy. One clinic visit only.")
        self.assertIn('clinic', entities['keywords'])
        # "patients" appears once — not a "frequently mentioned" term.
        self.assertNotIn('patients', entities['keywords'])

    def test_no_entities_in_empty_text(self):
        from services.documents.entities import extract_entities

        entities = extract_entities('')
        self.assertEqual(entities, {'numbers': [], 'dates': [], 'keywords': []})


class InsightsComputationTests(SimpleTestCase):
    def test_range_insight_for_numeric_column(self):
        from services.documents.insights import compute_insights

        datasets = [{
            'name': 'Sales', 'row_count': 3, 'duplicate_row_count': 0,
            'rows': [{'Revenue': 100}, {'Revenue': 300}, {'Revenue': 200}],
            'columns': [{
                'name': 'Revenue', 'display_name': 'Revenue', 'semantic_type': 'currency',
                'stats': {'min': 100, 'max': 300, 'mean': 200, 'std': 100, 'unique_count': 3, 'missing_count': 0},
            }],
        }]
        insights = compute_insights(datasets)
        self.assertTrue(any('100' in i and '300' in i for i in insights))

    def test_no_insight_when_column_is_constant(self):
        from services.documents.insights import compute_insights

        datasets = [{
            'name': 'Flat', 'row_count': 3, 'duplicate_row_count': 0,
            'rows': [{'Value': 5}, {'Value': 5}, {'Value': 5}],
            'columns': [{
                'name': 'Value', 'display_name': 'Value', 'semantic_type': 'number',
                'stats': {'min': 5, 'max': 5, 'mean': 5, 'std': 0, 'unique_count': 1, 'missing_count': 0},
            }],
        }]
        self.assertEqual(compute_insights(datasets), [])

    def test_no_insights_for_empty_dataset_list(self):
        from services.documents.insights import compute_insights

        self.assertEqual(compute_insights([]), [])


class DuplicateRowDetectionTests(TestCase):

    def test_duplicate_rows_are_counted(self):
        from services.pipeline import process_report_file

        csv_content = 'Name,City\nAda,Kigali\nBob,Musanze\nAda,Kigali\n'
        tmp_path = os.path.join(settings.BASE_DIR, 'uploads', 'tmp_dup_test.csv')
        with open(tmp_path, 'w', encoding='utf-8') as f:
            f.write(csv_content)
        try:
            result = process_report_file(tmp_path, 'tmp_dup_test.csv')
            dataset = result['data']['datasets'][0]
            self.assertEqual(dataset['duplicate_row_count'], 1)
        finally:
            os.remove(tmp_path)


class DocumentClassificationTests(SimpleTestCase):
    def test_classifies_sales_dataset(self):
        datasets = [{'columns': [
            {'name': 'Region', 'display_name': 'Region'},
            {'name': 'Product', 'display_name': 'Product'},
            {'name': 'Revenue', 'display_name': 'Revenue'},
            {'name': 'Units Sold', 'display_name': 'Units Sold'},
        ]}]
        label, confidence = classify_document(datasets)
        self.assertEqual(label, 'Sales')
        self.assertIn(confidence, ('high', 'medium'))

    def test_falls_back_to_general_dataset_when_unrecognized(self):
        datasets = [{'columns': [{'name': 'Foo', 'display_name': 'Foo'}, {'name': 'Bar', 'display_name': 'Bar'}]}]
        label, confidence = classify_document(datasets)
        self.assertEqual(label, 'General dataset')
        self.assertEqual(confidence, 'low')


class GenericCsvUploadTests(AuthenticatedTestCase):
    def setUp(self):
        super().setUp()
        with open(SALES_CSV, 'rb') as f:
            self.response = self.client.post('/api/reports/upload/', {'file': f}, format='multipart')

    def test_upload_succeeds_and_is_routed_generically(self):
        self.assertEqual(self.response.status_code, 201)
        data = self.response.json()
        self.assertEqual(data['status'], Report.Status.COMPLETED)
        self.assertEqual(data['source_type'], 'csv')
        self.assertEqual(data['extracted_data']['kind'], 'generic_document')

    def test_dataset_and_columns_extracted_correctly(self):
        document = self.response.json()['extracted_data']
        self.assertEqual(len(document['datasets']), 1)
        dataset = document['datasets'][0]
        self.assertEqual(dataset['row_count'], 12)
        semantic_types = {c['name']: c['semantic_type'] for c in dataset['columns']}
        self.assertEqual(semantic_types['Date'], 'date')
        self.assertEqual(semantic_types['Region'], 'category')
        self.assertEqual(semantic_types['Revenue'], 'currency')
        self.assertEqual(semantic_types['Discount %'], 'percentage')

    def test_document_classified_as_sales(self):
        document = self.response.json()['extracted_data']
        self.assertEqual(document['document_type'], 'Sales')

    def test_metrics_are_real_aggregates_not_invented(self):
        document = self.response.json()['extracted_data']
        revenue_metric = next(m for m in document['metrics'] if m['column'] == 'Revenue')
        self.assertAlmostEqual(revenue_metric['value'], 40875000.0)

    def test_download_returns_workbook_with_dataset_sheet(self):
        
        report_id = self.response.json()['id']
        download = self.client.get(f'/api/reports/{report_id}/download/')
        self.assertEqual(download.status_code, 200)
        workbook = openpyxl.load_workbook(__import__('io').BytesIO(b''.join(download.streaming_content)))
        self.assertIn('OVERVIEW', workbook.sheetnames)
        self.assertIn('Data', workbook.sheetnames)
        rows = list(workbook['Data'].iter_rows(min_row=2, values_only=True))
        self.assertEqual(len(rows), 12)


class GenericExcelUploadTests(AuthenticatedTestCase):
    def test_multi_sheet_excel_produces_multiple_datasets(self):
        with open(HR_XLSX, 'rb') as f:
            response = self.client.post('/api/reports/upload/', {'file': f}, format='multipart')
        self.assertEqual(response.status_code, 201)
        document = response.json()['extracted_data']
        self.assertEqual(document['kind'], 'generic_document')
        dataset_names = {d['name'] for d in document['datasets']}
        self.assertEqual(dataset_names, {'Employees', 'Departments'})
        self.assertEqual(document['document_type'], 'HR')


class LegacyXlsUploadTests(AuthenticatedTestCase):
    """Legacy binary .xls workbooks previously fell through EXCEL_EXTENSIONS
    into openpyxl, which cannot open them at all — this exercises the xlrd
    fallback added in services.extraction.excel_extractor."""

    def test_legacy_xls_workbook_is_extracted_correctly(self):
        xls_path = os.path.join(SAMPLE_DIR, 'generic_inventory_sample.xls')
        with open(xls_path, 'rb') as f:
            response = self.client.post('/api/reports/upload/', {'file': f}, format='multipart')
        self.assertEqual(response.status_code, 201)
        data = response.json()
        self.assertEqual(data['source_type'], 'excel')
        document = data['extracted_data']
        self.assertEqual(document['kind'], 'generic_document')
        dataset = document['datasets'][0]
        self.assertEqual(dataset['row_count'], 4)
        cable_reel = next(r for r in dataset['rows'] if r['Item'] == 'Cable Reel')
        self.assertEqual(cable_reel['Quantity'], 100.0)


class GenericDocxUploadTests(AuthenticatedTestCase):

    def setUp(self):
        super().setUp()
        with open(OPERATIONS_DOCX, 'rb') as f:
            self.response = self.client.post('/api/reports/upload/', {'file': f}, format='multipart')

    def test_upload_succeeds(self):
        self.assertEqual(self.response.status_code, 201)
        data = self.response.json()
        self.assertEqual(data['source_type'], 'docx')
        self.assertEqual(data['extracted_data']['kind'], 'generic_document')

    def test_both_tables_extracted_as_separate_datasets(self):
        document = self.response.json()['extracted_data']
        self.assertEqual(len(document['datasets']), 2)
        row_counts = sorted(d['row_count'] for d in document['datasets'])
        self.assertEqual(row_counts, [3, 3])

    def test_narrative_text_preserved_as_sections_with_headings(self):
        document = self.response.json()['extracted_data']
        headings = [s['title'] for s in document['sections']]
        self.assertIn('Quarterly Operations Report', headings)
        self.assertIn('Outlook', headings)

    def test_table_values_are_not_invented(self):
        document = self.response.json()['extracted_data']
        throughput = next(d for d in document['datasets'] if 'Shipments' in {c['name'] for c in d['columns']})
        kigali = next(r for r in throughput['rows'] if r['Center'] == 'Kigali Hub')
        self.assertEqual(kigali['Shipments'], '4820')
        self.assertEqual(kigali['On-Time %'], '96.5')


class GenericTxtUploadTests(AuthenticatedTestCase):

    def setUp(self):
        super().setUp()
        with open(NOTES_TXT, 'rb') as f:
            self.response = self.client.post('/api/reports/upload/', {'file': f}, format='multipart')

    def test_upload_succeeds(self):
        self.assertEqual(self.response.status_code, 201)
        data = self.response.json()
        self.assertEqual(data['source_type'], 'txt')
        self.assertEqual(data['extracted_data']['kind'], 'generic_document')

    def test_embedded_table_detected(self):
        document = self.response.json()['extracted_data']
        self.assertEqual(len(document['datasets']), 1)
        dataset = document['datasets'][0]
        self.assertEqual(dataset['row_count'], 5)
        semantic_types = {c['name']: c['semantic_type'] for c in dataset['columns']}
        self.assertEqual(semantic_types['Cost (USD)'], 'currency')

    def test_headings_recognised_around_the_table(self):
        document = self.response.json()['extracted_data']
        headings = [s['title'] for s in document['sections']]
        self.assertIn('Monthly Fuel Consumption Log', headings)
        self.assertIn('Notes', headings)


class ProseOnlyDocumentTests(AuthenticatedTestCase):
    def setUp(self):
        super().setUp()
        with open(PROSE_ONLY_TXT, 'rb') as f:
            self.response = self.client.post('/api/reports/upload/', {'file': f}, format='multipart')

    def test_upload_succeeds_with_no_datasets(self):
        self.assertEqual(self.response.status_code, 201)
        self.assertEqual(self.response.json()['status'], Report.Status.COMPLETED)
        document = self.response.json()['extracted_data']
        self.assertEqual(document['datasets'], [])
        self.assertGreater(len(document['sections']), 0)

    def test_entities_extracted_from_prose(self):
        document = self.response.json()['extracted_data']
        entities = document['entities']
        self.assertIn('88%', entities['numbers'])
        self.assertIn('$2.4 million', entities['numbers'])
        self.assertTrue(any('2026' in d for d in entities['dates']))
        self.assertIn('program', entities['keywords'])

    def test_no_insights_or_charts_invented_without_data(self):
        document = self.response.json()['extracted_data']
        self.assertEqual(document['insights'], [])
        self.assertEqual(document['charts'], [])

    def test_download_returns_a_workbook_even_with_no_datasets(self):
        report_id = self.response.json()['id']
        download = self.client.get(f'/api/reports/{report_id}/download/')
        self.assertEqual(download.status_code, 200)


class GenericJsonUploadTests(AuthenticatedTestCase):
    def setUp(self):
        super().setUp()
        with open(ORDERS_JSON, 'rb') as f:
            self.response = self.client.post('/api/reports/upload/', {'file': f}, format='multipart')

    def test_upload_succeeds_and_is_routed_generically(self):
        self.assertEqual(self.response.status_code, 201)
        data = self.response.json()
        self.assertEqual(data['status'], Report.Status.COMPLETED)
        self.assertEqual(data['source_type'], 'json')
        self.assertEqual(data['extracted_data']['kind'], 'generic_document')

    def test_top_level_array_field_becomes_the_primary_dataset(self):
        document = self.response.json()['extracted_data']
        primary = next(d for d in document['datasets'] if d['name'] == 'Data')
        self.assertEqual(primary['row_count'], 4)
        semantic_types = {c['name']: c['semantic_type'] for c in primary['columns']}
        self.assertEqual(semantic_types['order_id'], 'identifier')
        self.assertEqual(semantic_types['total'], 'currency')
        self.assertEqual(semantic_types['placed_on'], 'date')

    def test_nested_array_of_objects_becomes_its_own_dataset(self):
        document = self.response.json()['extracted_data']
        nested = next(d for d in document['datasets'] if d['name'] == 'line_items')
        self.assertEqual(nested['row_count'], 7)
        column_names = {c['name'] for c in nested['columns']}
        self.assertEqual(column_names, {'sku', 'qty', 'unit_price'})

    def test_no_values_are_invented(self):
        document = self.response.json()['extracted_data']
        primary = next(d for d in document['datasets'] if d['name'] == 'Data')
        first_order = next(r for r in primary['rows'] if r['order_id'] == 'ORD-1002')
        self.assertEqual(first_order['customer'], 'Jean Bosco')
        self.assertEqual(first_order['total'], 76000)

    def test_rows_carry_their_source_row_for_traceability(self):
        document = self.response.json()['extracted_data']
        primary = next(d for d in document['datasets'] if d['name'] == 'Data')
        source_rows = sorted(r['_source_row'] for r in primary['rows'])
        self.assertEqual(source_rows, [2, 3, 4, 5])

    def test_download_returns_workbook_with_dictionary_and_metadata_sheets(self):
        report_id = self.response.json()['id']
        download = self.client.get(f'/api/reports/{report_id}/download/')
        self.assertEqual(download.status_code, 200)
        workbook = openpyxl.load_workbook(__import__('io').BytesIO(b''.join(download.streaming_content)))
        self.assertIn('DATA DICTIONARY', workbook.sheetnames)
        self.assertIn('SOURCE METADATA', workbook.sheetnames)
        self.assertIn('Data', workbook.sheetnames)
        header = [c.value for c in next(workbook['Data'].iter_rows(min_row=1, max_row=1))]
        self.assertEqual(header[0], 'Source Row')

    def test_json_export_returns_same_data_as_dashboard(self):
        report_id = self.response.json()['id']
        api_data = self.response.json()['extracted_data']
        export = self.client.get(f'/api/reports/{report_id}/download/json/')
        self.assertEqual(export.status_code, 200)
        payload = export.json()
        self.assertEqual(payload['data'], api_data)
        self.assertEqual(payload['source_type'], 'json')


class GenericJsonLinesUploadTests(AuthenticatedTestCase):
    def test_jsonl_file_is_parsed_one_record_per_line(self):
        from django.core.files.uploadedfile import SimpleUploadedFile

        content = (
            b'{"name": "Ada", "team": "Engineering", "score": 91}\n'
            b'{"name": "Bosco", "team": "Sales", "score": 78}\n'
            b'{"name": "Grace", "team": "Finance", "score": 85}\n'
        )
        upload = SimpleUploadedFile('scores.jsonl', content, content_type='application/x-ndjson')
        response = self.client.post('/api/reports/upload/', {'file': upload}, format='multipart')
        self.assertEqual(response.status_code, 201)
        document = response.json()['extracted_data']
        self.assertEqual(document['source_type'], 'json')
        dataset = document['datasets'][0]
        self.assertEqual(dataset['row_count'], 3)
        names = {r['name'] for r in dataset['rows']}
        self.assertEqual(names, {'Ada', 'Bosco', 'Grace'})

    def test_invalid_json_fails_clearly(self):
        from django.core.files.uploadedfile import SimpleUploadedFile

        upload = SimpleUploadedFile('broken.json', b'{not valid json', content_type='application/json')
        response = self.client.post('/api/reports/upload/', {'file': upload}, format='multipart')
        self.assertEqual(response.status_code, 422)
        self.assertIn('Invalid JSON', response.json()['error_message'])


class PdfWithFiguresAndInsightsTests(AuthenticatedTestCase):
    def setUp(self):
        super().setUp()
        with open(OPERATIONS_PDF, 'rb') as f:
            self.response = self.client.post('/api/reports/upload/', {'file': f}, format='multipart')

    def test_upload_succeeds(self):
        self.assertEqual(self.response.status_code, 201)
        data = self.response.json()
        self.assertEqual(data['source_type'], 'pdf')
        self.assertEqual(data['extracted_data']['kind'], 'generic_document')

    def test_table_extracted_correctly(self):
        document = self.response.json()['extracted_data']
        self.assertEqual(len(document['datasets']), 1)
        dataset = document['datasets'][0]
        self.assertEqual(dataset['row_count'], 3)
        kigali = next(r for r in dataset['rows'] if r['Center'] == 'Kigali Hub')
        self.assertEqual(kigali['Shipments'], '4820')

    def test_embedded_figure_extracted(self):
        document = self.response.json()['extracted_data']
        self.assertEqual(len(document['figures']), 1)
        figure = document['figures'][0]
        self.assertEqual(figure['source'], 'Page 1')
        self.assertTrue(figure['thumbnail'].startswith('data:image/png;base64,'))

    def test_insights_are_real_aggregates(self):
        document = self.response.json()['extracted_data']
        insight_text = ' '.join(document['insights'])
        self.assertIn('4,820', insight_text)
        self.assertIn('Kigali Hub', insight_text)

    def test_prose_section_preserved(self):
        document = self.response.json()['extracted_data']
        self.assertTrue(any('Regional Distribution Center Review' == s['title'] for s in document['sections']))


class JsonEnvelopeSelectionTests(SimpleTestCase):
    """Unit-level tests of json_extractor's envelope-selection rule — see
    the module docstring: a known key wins deterministically, an ambiguous
    set of arrays becomes multiple datasets with a warning, never a silent
    "pick the longest one" guess."""

    def test_known_envelope_key_is_picked_deterministically(self):
        from services.extraction.json_extractor import _select_root_tables

        payload = {'meta': {'page': 1}, 'results': [{'a': 1}, {'a': 2}, {'a': 3}], 'data': [{'a': 1}]}
        tables, warnings = _select_root_tables(payload)
        # 'data' precedes 'results' in the fixed priority order, and wins
        # even though 'results' has more items -- never "whichever is longest".
        self.assertEqual(list(tables.keys()), [''])
        self.assertEqual(len(tables['']), 1)
        self.assertEqual(warnings, [])

    def test_ambiguous_top_level_arrays_become_separate_datasets_with_a_warning(self):
        from services.extraction.json_extractor import _select_root_tables

        payload = {'orders': [{'id': 1}, {'id': 2}], 'customers': [{'id': 1}]}
        tables, warnings = _select_root_tables(payload)
        self.assertEqual(set(tables.keys()), {'orders', 'customers'})
        self.assertEqual(len(warnings), 1)
        self.assertIn('orders', warnings[0])
        self.assertIn('customers', warnings[0])

    def test_single_unrecognized_array_field_is_unambiguous(self):
        from services.extraction.json_extractor import _select_root_tables

        payload = {'meta': {'page': 1}, 'orders': [{'id': 1}]}
        tables, warnings = _select_root_tables(payload)
        self.assertEqual(list(tables.keys()), [''])
        self.assertEqual(warnings, [])


class JsonProvenanceTests(AuthenticatedTestCase):
    """A nested JSON array's rows must always be traceable back to their
    exact position in the source document and to the parent row that
    produced them."""

    def setUp(self):
        super().setUp()
        with open(ORDERS_JSON, 'rb') as f:
            self.response = self.client.post('/api/reports/upload/', {'file': f}, format='multipart')

    def test_top_level_rows_carry_a_json_path(self):
        document = self.response.json()['extracted_data']
        primary = next(d for d in document['datasets'] if d['name'] == 'Data')
        paths = sorted(r['_json_path'] for r in primary['rows'])
        # "orders" is the *only* top-level array field, so it's unambiguous
        # and becomes the primary/unnamed table -- its JSONPath is relative
        # to the document root, not prefixed with the field name.
        self.assertEqual(paths, ['$[0]', '$[1]', '$[2]', '$[3]'])

    def test_nested_rows_reference_their_parent_order(self):
        document = self.response.json()['extracted_data']
        nested = next(d for d in document['datasets'] if d['name'] == 'line_items')
        parent_refs = {r['_parent_ref'] for r in nested['rows']}
        # order_id is the parent's own identifier field, used in preference
        # to a positional fallback.
        self.assertTrue(parent_refs.issubset({'ORD-1001', 'ORD-1002', 'ORD-1003', 'ORD-1004'}))
        ord_1002_items = [r for r in nested['rows'] if r['_parent_ref'] == 'ORD-1002']
        self.assertEqual(len(ord_1002_items), 1)
        self.assertEqual(ord_1002_items[0]['sku'], 'BP-300')


class PdfTableMergeTests(SimpleTestCase):
    """Unit-level tests of the multi-page table merge heuristic — merges
    only when a table is the first on the very next page with an identical
    header, otherwise leaves tables separate (the safe default)."""

    def test_identical_header_on_next_page_is_merged(self):
        from services.documents.generic_parser import _merge_continued_pdf_tables

        page1 = [[['Name', 'Amount'], ['Ada', '100']]]
        page2 = [[['Name', 'Amount'], ['Bob', '200']]]
        result = _merge_continued_pdf_tables([page1, page2])
        self.assertEqual(len(result), 1)
        source, table = result[0]
        self.assertIn('merged across page break', source)
        self.assertEqual(table, [['Name', 'Amount'], ['Ada', '100'], ['Bob', '200']])

    def test_different_header_on_next_page_is_not_merged(self):
        from services.documents.generic_parser import _merge_continued_pdf_tables

        page1 = [[['Name', 'Amount'], ['Ada', '100']]]
        page2 = [[['Region', 'Total'], ['East', '500']]]
        result = _merge_continued_pdf_tables([page1, page2])
        self.assertEqual(len(result), 2)

    def test_table_not_first_on_next_page_is_not_merged(self):
        from services.documents.generic_parser import _merge_continued_pdf_tables

        page1 = [[['Name', 'Amount'], ['Ada', '100']]]
        page2 = [[['Other', 'Table'], ['x', 'y']], [['Name', 'Amount'], ['Bob', '200']]]
        result = _merge_continued_pdf_tables([page1, page2])
        # The matching-header table is second on page 2, not first -- no merge.
        self.assertEqual(len(result), 3)


class ScannedPdfTests(AuthenticatedTestCase):
    """A PDF with no extractable text layer at all (the hallmark of a
    scanned image) must fail clearly, and can be recovered exactly once an
    OCR engine is plugged into the OCR_ENGINE extension point."""

    BLANK_PDF = os.path.join(SAMPLE_DIR, 'generic_blank_no_text_sample.pdf')

    def test_scanned_pdf_fails_with_a_specific_message_not_a_generic_one(self):
        with open(self.BLANK_PDF, 'rb') as f:
            response = self.client.post('/api/reports/upload/', {'file': f}, format='multipart')
        self.assertEqual(response.status_code, 422)
        message = response.json()['error_message']
        self.assertIn('scanned image', message)
        self.assertIn('OCR', message)

    def test_ocr_extension_point_recovers_text_and_flags_for_review(self):
        from services.extraction import pdf_extractor

        original_engine = pdf_extractor.OCR_ENGINE
        pdf_extractor.OCR_ENGINE = lambda path: 'Recovered by OCR.\n\nThis is a long enough paragraph of OCR output.'
        try:
            with open(self.BLANK_PDF, 'rb') as f:
                response = self.client.post('/api/reports/upload/', {'file': f}, format='multipart')
        finally:
            pdf_extractor.OCR_ENGINE = original_engine

        self.assertEqual(response.status_code, 201)
        data = response.json()
        self.assertTrue(any('OCR' in w for w in data['warnings']))


class UploadSecurityTests(AuthenticatedTestCase):
    def test_duplicate_upload_is_recognized_instead_of_reprocessed(self):
        with open(SALES_CSV, 'rb') as f:
            first = self.client.post('/api/reports/upload/', {'file': f}, format='multipart')
        self.assertEqual(first.status_code, 201)

        with open(SALES_CSV, 'rb') as f:
            second = self.client.post('/api/reports/upload/', {'file': f}, format='multipart')
        self.assertEqual(second.status_code, 200)
        self.assertEqual(second.json()['id'], first.json()['id'])

    def test_content_not_matching_claimed_extension_is_rejected(self):
        from django.core.files.uploadedfile import SimpleUploadedFile

        # A plain-text file renamed to .pdf: no "%PDF-" signature.
        upload = SimpleUploadedFile('fake.pdf', b'This is not really a PDF file at all.', content_type='application/pdf')
        response = self.client.post('/api/reports/upload/', {'file': upload}, format='multipart')
        self.assertEqual(response.status_code, 400)
        self.assertIn('content', response.json()['detail'].lower())

    def test_path_traversal_filename_is_sanitized(self):
        with open(SALES_CSV, 'rb') as f:
            content = f.read()
        from django.core.files.uploadedfile import SimpleUploadedFile

        upload = SimpleUploadedFile('../../etc/evil_sales.csv', content, content_type='text/csv')
        response = self.client.post('/api/reports/upload/', {'file': upload}, format='multipart')
        self.assertEqual(response.status_code, 201)
        self.assertEqual(response.json()['original_filename'], 'evil_sales.csv')

    def test_decompression_bomb_is_rejected(self):
        import zipfile
        from django.core.files.uploadedfile import SimpleUploadedFile

        scratch_path = os.path.join(settings.BASE_DIR, 'uploads', 'tmp_zipbomb.xlsx')
        with zipfile.ZipFile(scratch_path, 'w', compression=zipfile.ZIP_DEFLATED) as zf:
            zf.writestr('huge.txt', b'0' * (50 * 1024 * 1024))
        try:
            with open(scratch_path, 'rb') as f:
                data = f.read()
            upload = SimpleUploadedFile(
                'bomb.xlsx', data,
                content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            )
            response = self.client.post('/api/reports/upload/', {'file': upload}, format='multipart')
            self.assertEqual(response.status_code, 422)
            self.assertIn('decompress', response.json()['error_message'].lower())
        finally:
            os.remove(scratch_path)


class ExcelHardeningTests(AuthenticatedTestCase):
    """Merged cells, formulas with no cached value, and hidden sheets are
    all real shapes production spreadsheets take — each must be handled
    without silently mistaking one thing for another (a merged cell's blank
    continuation for "no data", a stale formula for "genuinely zero", a
    hidden sheet for "not part of the workbook")."""

    def setUp(self):
        super().setUp()
        import io

        import openpyxl

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Regions'
        ws['A1'] = 'Region'
        ws['B1'] = 'Sub-region'
        ws['C1'] = 'Population'
        ws['D1'] = 'Computed'
        ws['A2'] = 'East Africa'
        ws.merge_cells('A2:A3')
        ws['B2'] = 'North'
        ws['C2'] = 100
        ws['D2'] = '=C2*2'  # never calculated -> no cached value
        ws['B3'] = 'South'
        ws['C3'] = 200
        ws['A4'] = 'West Africa'
        ws['B4'] = 'Central'
        ws['C4'] = 50

        hidden = wb.create_sheet('Notes')
        hidden.sheet_state = 'hidden'
        hidden['A1'] = 'Col1'
        hidden['B1'] = 'Col2'
        hidden['A2'] = 'a'
        hidden['B2'] = 'b'

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)

        from django.core.files.uploadedfile import SimpleUploadedFile
        upload = SimpleUploadedFile(
            'regions.xlsx', buf.read(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        self.response = self.client.post('/api/reports/upload/', {'file': upload}, format='multipart')

    def test_upload_succeeds(self):
        self.assertEqual(self.response.status_code, 201)

    def test_merged_cell_value_propagates_to_covered_rows(self):
        document = self.response.json()['extracted_data']
        regions = next(d for d in document['datasets'] if d['name'] == 'Regions')
        south_row = next(r for r in regions['rows'] if r['Sub-region'] == 'South')
        self.assertEqual(south_row['Region'], 'East Africa')

    def test_formula_without_cached_value_raises_a_warning(self):
        data = self.response.json()
        self.assertTrue(any('cached value' in w for w in data['warnings']))

    def test_hidden_sheet_is_still_processed_but_flagged_as_info(self):
        document = self.response.json()['extracted_data']
        dataset_names = {d['name'] for d in document['datasets']}
        self.assertIn('Notes', dataset_names)  # never silently skipped

        issues = document['validation_issues']
        hidden_issues = [i for i in issues if 'hidden' in i['message'].lower()]
        self.assertEqual(len(hidden_issues), 1)
        self.assertEqual(hidden_issues[0]['severity'], 'info')
        # info-severity issues stay out of the main warnings panel.
        self.assertFalse(any('hidden' in w.lower() for w in self.response.json()['warnings']))


class ValidationIssuesStructureTests(AuthenticatedTestCase):
    def test_every_issue_has_the_documented_shape(self):
        with open(SALES_CSV, 'rb') as f:
            response = self.client.post('/api/reports/upload/', {'file': f}, format='multipart')
        document = response.json()['extracted_data']
        self.assertIn('validation_issues', document)
        for issue in document['validation_issues']:
            self.assertIn(issue['severity'], ('info', 'warning', 'error'))
            for key in ('rule_id', 'severity', 'message', 'dataset', 'column', 'row', 'source', 'value', 'suggested_action'):
                self.assertIn(key, issue)


class ReconciliationAndExportHardeningTests(AuthenticatedTestCase):
    def setUp(self):
        super().setUp()
        with open(SALES_CSV, 'rb') as f:
            self.response = self.client.post('/api/reports/upload/', {'file': f}, format='multipart')

    def _download_workbook(self):
        import io as _io

        report_id = self.response.json()['id']
        download = self.client.get(f'/api/reports/{report_id}/download/')
        return openpyxl.load_workbook(_io.BytesIO(b''.join(download.streaming_content)))

    def test_reconciliation_sheet_accounts_for_every_row(self):
        workbook = self._download_workbook()
        self.assertIn('RECONCILIATION', workbook.sheetnames)
        ws = workbook['RECONCILIATION']
        row = next(r for r in ws.iter_rows(min_row=3, values_only=True) if r[0] == 'Data')
        # 12 data rows + 1 header row, no blanks, no duplicates.
        self.assertEqual(row[1], 13)
        self.assertEqual(row[2], 0)
        self.assertEqual(row[3], 12)

    def test_dataset_sheet_is_a_real_excel_table(self):
        workbook = self._download_workbook()
        ws = workbook['Data']
        self.assertGreaterEqual(len(ws.tables), 1)

    def test_workbook_reopens_cleanly_with_openpyxl(self):
        # Implicit in every assertion above (load_workbook would raise on a
        # malformed file) -- asserted explicitly here too for clarity.
        workbook = self._download_workbook()
        self.assertIn('OVERVIEW', workbook.sheetnames)


class FormulaInjectionExportTests(AuthenticatedTestCase):
    def test_formula_looking_category_value_is_not_evaluated_on_export(self):
        from django.core.files.uploadedfile import SimpleUploadedFile

        csv_content = (
            b'Name,Note\n'
            b'Ada,Fine\n'
            b'=cmd|/C calc!A1,Suspicious\n'
            b'Bob,Fine\n'
        )
        upload = SimpleUploadedFile('notes.csv', csv_content, content_type='text/csv')
        response = self.client.post('/api/reports/upload/', {'file': upload}, format='multipart')
        self.assertEqual(response.status_code, 201)

        report_id = response.json()['id']
        download = self.client.get(f'/api/reports/{report_id}/download/')
        import io as _io
        workbook = openpyxl.load_workbook(_io.BytesIO(b''.join(download.streaming_content)))
        ws = workbook['Data']
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith('='):
                    self.assertEqual(cell.data_type, 's', f'{cell.coordinate} was stored as a formula, not text')
