import io
import os

import openpyxl
from django.conf import settings
from django.contrib.auth import get_user_model
from django.test import SimpleTestCase, TestCase

from reports.models import Report
from services.normalization.validate_data import _validate_bond_list

User = get_user_model()

SAMPLE_DIR = os.path.join(settings.BASE_DIR.parent, 'sample_data')
SAMPLE_PDF = os.path.join(SAMPLE_DIR, 'RSE_sample.pdf')
SAMPLE_XLSX = os.path.join(SAMPLE_DIR, 'RSE_sample.xlsx')


class AuthenticatedTestCase(TestCase):

    def setUp(self):
        self.user = User.objects.create_user(username='tester@example.com', email='tester@example.com', password='test-pass-123')
        self.client.login(username='tester@example.com', password='test-pass-123')
        super().setUp()


class ReportUploadPdfTests(AuthenticatedTestCase):
    def setUp(self):
        super().setUp()
        with open(SAMPLE_PDF, 'rb') as f:
            self.response = self.client.post('/api/reports/upload/', {'file': f}, format='multipart')

    def test_upload_succeeds(self):
        self.assertEqual(self.response.status_code, 201)
        self.assertEqual(self.response.json()['status'], Report.Status.COMPLETED)

    def test_extracted_sections_are_populated(self):
        data = self.response.json()['extracted_data']
        self.assertEqual(data['report_date'], '2026-09-07')
        self.assertEqual(len(data['equities']), 10)
        self.assertEqual(len(data['government_bonds']), 39)
        self.assertEqual(len(data['corporate_bonds']), 10)
        self.assertEqual(len(data['bond_trades']), 5)
        self.assertEqual(len(data['exchange_rates']), 6)
        index_names = {i['name'] for i in data['indices']}
        self.assertEqual(index_names, {'RSI', 'ALSI'})

    def test_bond_market_overview_is_populated(self):
        overview = self.response.json()['extracted_data']['market_overview']
        self.assertEqual(overview['bond_turnover'], 206685500)
        self.assertEqual(overview['bond_deals'], 5)
        self.assertEqual(overview['market_capitalization'], 6634919683716)

    def test_no_values_are_invented(self):
        data = self.response.json()['extracted_data']
        bok = next(e for e in data['equities'] if e['ticker'] == 'BOK')
        self.assertEqual(bok['closing'], 660)
        self.assertEqual(bok['volume'], 21600)
        self.assertEqual(bok['value'], 14256000)

    def test_market_overview_matches_source_narrative(self):
        overview = self.response.json()['extracted_data']['market_overview']
        self.assertEqual(overview['equity_turnover'], 14307500)
        self.assertEqual(overview['shares_traded'], 21700)
        self.assertEqual(overview['equity_deals'], 4)
        self.assertEqual(overview['bond_turnover'], 206685500)
        self.assertEqual(overview['bond_deals'], 5)
        self.assertEqual(overview['market_capitalization'], 6634919683716)

    def test_repo_market_activity_matches_source(self):
        overview = self.response.json()['extracted_data']['market_overview']
        self.assertEqual(overview['repo_deals'], 3)
        self.assertEqual(overview['repo_turnover'], 10_000_000_000)
        self.assertEqual(overview['repo_tenor'], '7-day')
        self.assertEqual(overview['repo_rate'], 8.75)

    def test_indices_match_source(self):
        indices = {i['name']: i for i in self.response.json()['extracted_data']['indices']}
        self.assertEqual(indices['RSI']['closing'], 214.08)
        self.assertEqual(indices['RSI']['previous'], 214.08)
        self.assertEqual(indices['ALSI']['closing'], 259.21)
        self.assertEqual(indices['ALSI']['previous'], 259.21)

    def test_all_equities_match_source(self):
        expected = {
            'BOK': (660, 21600, 14256000),
            'BLR': (515, 100, 51500),
            'NMG': (1200, 0, 0),
            'KCB': (500, 0, 0),
            'USL': (104, 0, 0),
            'EQTY': (900, 0, 0),
            'IMR': (100, 0, 0),
            'RHB': (526, 0, 0),
            'CMR': (165, 0, 0),
            'MTNR': (136, 0, 0),
        }
        equities = {e['ticker']: e for e in self.response.json()['extracted_data']['equities']}
        self.assertEqual(set(equities), set(expected))
        for ticker, (closing, volume, value) in expected.items():
            self.assertEqual(equities[ticker]['closing'], closing, ticker)
            self.assertEqual(equities[ticker]['volume'], volume, ticker)
            self.assertEqual(equities[ticker]['value'], value, ticker)

    def test_exchange_rates_match_source(self):
        expected = {
            'USD': (1466.39, 1476.39, 1471.39),
            'KES': (11.32, 11.40, 11.36),
            'UGS': (0.38, 0.39, 0.38),
            'BIF': (0.49, 0.49, 0.49),
            'TZS': (0.55, 0.55, 0.55),
            'ZAR': (91.77, 92.40, 92.08),
        }
        rates = {r['currency']: r for r in self.response.json()['extracted_data']['exchange_rates']}
        self.assertEqual(set(rates), set(expected))
        for currency, (buying, selling, average) in expected.items():
            self.assertEqual(rates[currency]['buying'], buying, currency)
            self.assertEqual(rates[currency]['selling'], selling, currency)
            self.assertEqual(rates[currency]['average'], average, currency)

    def test_bond_trades_match_source(self):
        expected = {
            'FXD2/2020/15Yrs (Re-opened)': (52000000, 100.85, 100.8),
            'FXD3/2021/15Yrs (Re-opened)': (50000000, 102, 100.7),
            'FXD1/2023/20Yrs (Re-opened)': (50000000, 103.5, 101.9),
            'FXD1/2024/20Yrs (Re-opened)': (1500000, 101.4, 101.3),
            'FXD8/2024/20Yrs (Re-opened)': (50000000, 102.7, 102.9),
        }
        trades = {t['bond']: t for t in self.response.json()['extracted_data']['bond_trades']}
        self.assertEqual(set(trades), set(expected))
        for bond, (volume, previous, closing) in expected.items():
            self.assertEqual(trades[bond]['volume'], volume, bond)
            self.assertEqual(trades[bond]['previous'], previous, bond)
            self.assertEqual(trades[bond]['closing'], closing, bond)
            self.assertEqual(trades[bond]['category'], 'TREASURY', bond)

    def test_mgmrw_corporate_bonds_are_both_preserved_without_false_warning(self):
        
        data = self.response.json()['extracted_data']
        mgm_bonds = [b for b in data['corporate_bonds'] if b['isin'] == 'MGMRW']
        self.assertEqual(len(mgm_bonds), 2)

        by_security = {b['security']: b for b in mgm_bonds}
        self.assertEqual(by_security['MGM/2024/5YRS']['maturity_date'], '2029-08-28')
        self.assertEqual(by_security['MGM/2024/5YRS']['coupon_rate'], 15.0)
        self.assertEqual(by_security['MGM/2024/5YRS']['closing_price'], 100.0)
        self.assertEqual(by_security['MGM/2025/5YRS']['maturity_date'], '2030-08-02')
        self.assertEqual(by_security['MGM/2025/5YRS']['coupon_rate'], 14.0)
        self.assertEqual(by_security['MGM/2025/5YRS']['closing_price'], 104.27)

        for warning in self.response.json()['warnings']:
            self.assertNotIn('MGMRW', warning, f'False duplicate warning regressed: {warning!r}')

    def test_all_corporate_bonds_present(self):
        expected_securities = {
            'ECTL/2021/10Yrs', 'BSLB1/2023/7YRS', 'MGM/2024/5YRS', 'BSLB2/2024/7YRS',
            'PRE/2024/7YRS', 'IFC/2025/8YRS', 'AMS/2025/5YRS', 'MGM/2025/5YRS',
            'ECTL/2025/7YRS', 'BSLB3/2026/10YRS',
        }
        data = self.response.json()['extracted_data']
        actual_securities = {b['security'] for b in data['corporate_bonds']}
        self.assertEqual(actual_securities, expected_securities)

    def test_download_returns_workbook(self):
        report_id = self.response.json()['id']
        download = self.client.get(f'/api/reports/{report_id}/download/')
        self.assertEqual(download.status_code, 200)
        self.assertEqual(
            download['Content-Type'],
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )

    def test_downloaded_workbook_matches_dashboard_data(self):
        report_id = self.response.json()['id']
        api_data = self.response.json()['extracted_data']

        download = self.client.get(f'/api/reports/{report_id}/download/')
        workbook = openpyxl.load_workbook(io.BytesIO(b''.join(download.streaming_content)))

        self.assertEqual(
            set(workbook.sheetnames),
            {'MARKET SUMMARY', 'STOCK', 'MARKET STATS', 'BONDS', 'BONDS TRADES', 'EXCHANGE RATE', 'INDICES'},
        )

        bonds_ws = workbook['BONDS']
        bond_rows = list(bonds_ws.iter_rows(min_row=2, values_only=True))
        self.assertEqual(len(bond_rows), len(api_data['government_bonds']) + len(api_data['corporate_bonds']))
        mgm_rows = [r for r in bond_rows if r[0] == 'MGMRW']
        self.assertEqual(len(mgm_rows), 2)
        mgm_securities = {r[2] for r in mgm_rows}
        self.assertEqual(mgm_securities, {'MGM/2024/5YRS', 'MGM/2025/5YRS'})

        stock_ws = workbook['STOCK']
        stock_rows = list(stock_ws.iter_rows(min_row=2, values_only=True))
        self.assertEqual(len(stock_rows), len(api_data['equities']))
        bok_row = next(r for r in stock_rows if r[1] == 'BOK')
        self.assertEqual(bok_row[6], 660)  
        self.assertEqual(bok_row[9], 21600) 

        rates_ws = workbook['EXCHANGE RATE']
        rate_rows = {r[0]: r for r in rates_ws.iter_rows(min_row=2, values_only=True)}
        self.assertEqual(rate_rows['USD'][1:3], (1466.39, 1476.39))


class ReportUploadExcelTests(AuthenticatedTestCase):
    def test_upload_excel_succeeds(self):
        with open(SAMPLE_XLSX, 'rb') as f:
            response = self.client.post('/api/reports/upload/', {'file': f}, format='multipart')
        self.assertEqual(response.status_code, 201)
        data = response.json()['extracted_data']
        self.assertEqual(len(data['equities']), 10)
        self.assertEqual(len(data['exchange_rates']), 6)
        self.assertEqual(len(data['bond_trades']), 5)


class ReportUploadValidationTests(AuthenticatedTestCase):
    def test_rejects_unsupported_file_type(self):
        fake_file = io.BytesIO(b'binary content')
        fake_file.name = 'notes.zip'
        response = self.client.post('/api/reports/upload/', {'file': fake_file}, format='multipart')
        self.assertEqual(response.status_code, 400)

    def test_prose_only_txt_still_succeeds_with_a_document_overview(self):
        fake_file = io.BytesIO(b'This is just a paragraph of prose with no table or delimited data in it at all.')
        fake_file.name = 'notes.txt'
        response = self.client.post('/api/reports/upload/', {'file': fake_file}, format='multipart')
        self.assertEqual(response.status_code, 201)
        data = response.json()
        self.assertEqual(data['status'], Report.Status.COMPLETED)
        document = data['extracted_data']
        self.assertEqual(document['datasets'], [])
        self.assertTrue(any('paragraph of prose' in s['content'] for s in document['sections']))

    def test_genuinely_empty_document_fails_clearly(self):
        fake_file = io.BytesIO(b'   \n\n   ')
        fake_file.name = 'blank.txt'
        response = self.client.post('/api/reports/upload/', {'file': fake_file}, format='multipart')
        self.assertEqual(response.status_code, 422)
        self.assertEqual(response.json()['status'], Report.Status.FAILED)

    def test_rejects_missing_file(self):
        response = self.client.post('/api/reports/upload/', {}, format='multipart')
        self.assertEqual(response.status_code, 400)

    def test_unparseable_document_reports_failure_clearly(self):
        fake_file = io.BytesIO(b'%PDF-1.4\n%%EOF')
        fake_file.name = 'empty.pdf'
        response = self.client.post('/api/reports/upload/', {'file': fake_file}, format='multipart')
        self.assertIn(response.status_code, (422, 500))
        self.assertEqual(response.json()['status'], Report.Status.FAILED)
        self.assertTrue(response.json()['error_message'])


class ReportDetailTests(AuthenticatedTestCase):
    def test_missing_report_returns_404(self):
        response = self.client.get('/api/reports/00000000-0000-0000-0000-000000000000/')
        self.assertEqual(response.status_code, 404)


class ReportAuthAndOwnershipTests(TestCase):
    def setUp(self):
        self.owner = User.objects.create_user(username='owner@example.com', email='owner@example.com', password='owner-pass-123')
        self.other = User.objects.create_user(username='other@example.com', email='other@example.com', password='other-pass-123')

    def test_upload_requires_authentication(self):
        with open(SAMPLE_PDF, 'rb') as f:
            response = self.client.post('/api/reports/upload/', {'file': f}, format='multipart')
        self.assertEqual(response.status_code, 403)

    def test_list_requires_authentication(self):
        self.assertEqual(self.client.get('/api/reports/').status_code, 403)

    def test_uploaded_report_is_scoped_to_its_owner(self):
        self.client.login(username='owner@example.com', password='owner-pass-123')
        with open(SAMPLE_PDF, 'rb') as f:
            upload = self.client.post('/api/reports/upload/', {'file': f}, format='multipart')
        report_id = upload.json()['id']

        listing = self.client.get('/api/reports/')
        self.assertEqual(listing.status_code, 200)
        self.assertEqual([r['id'] for r in listing.json()['results']], [report_id])
        self.assertEqual(self.client.get(f'/api/reports/{report_id}/').status_code, 200)

        self.client.logout()
        self.client.login(username='other@example.com', password='other-pass-123')
        self.assertEqual(self.client.get('/api/reports/').json()['results'], [])
        self.assertEqual(self.client.get(f'/api/reports/{report_id}/').status_code, 404)
        self.assertEqual(self.client.get(f'/api/reports/{report_id}/download/').status_code, 404)
        self.assertEqual(self.client.patch(f'/api/reports/{report_id}/', {'original_filename': 'x.pdf'}, content_type='application/json').status_code, 404)
        self.assertEqual(self.client.delete(f'/api/reports/{report_id}/').status_code, 404)


def _bond(isin, security, maturity, coupon, closing=100.0):
    return {
        'isin': isin, 'status': None, 'security': security, 'label': security,
        'maturity_date': maturity, 'coupon_rate': coupon, 'closing_price': closing,
        'previous_price': closing, 'bids': 0, 'offers': 0, 'bond_traded': 0,
        'category': 'CORPORATE', 'issue_date': None, 'yield_tm': None, 't_bond_no': None,
    }


class BondDuplicateValidationTests(SimpleTestCase):
    def test_same_issuer_code_different_series_is_not_a_duplicate(self):
        bonds = [
            _bond('MGMRW', 'MGM/2024/5YRS', '2029-08-28', 15.0, 100.0),
            _bond('MGMRW', 'MGM/2025/5YRS', '2030-08-02', 14.0, 104.27),
        ]
        warnings = []
        valid = _validate_bond_list(bonds, warnings, 'corporate')
        self.assertEqual(len(valid), 2)
        self.assertEqual(warnings, [])

    def test_truly_identical_bond_row_is_still_flagged(self):
        bonds = [
            _bond('RW000A182K48', 'FXD2/2016/15Yrs', '2031-05-09', 13.5, 103.0),
            _bond('RW000A182K48', 'FXD2/2016/15Yrs', '2031-05-09', 13.5, 103.0),
        ]
        warnings = []
        valid = _validate_bond_list(bonds, warnings, 'government')
        self.assertEqual(len(valid), 2) 
        self.assertEqual(len(warnings), 1)
        self.assertIn('RW000A182K48', warnings[0])
        self.assertIn('duplicate', warnings[0].lower())
