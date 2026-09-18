from __future__ import annotations

import os
from dataclasses import dataclass, field
from typing import Dict, List

import openpyxl

from services.normalization.clean_data import is_blank_row

from .zip_guard import check_zip_bomb

SheetRows = List[List]

# Above this size, the extra non-streaming pass used for merged-cell
# fill-forward and formula-without-cached-value detection is skipped —
# it requires loading the whole workbook (not just streaming rows), so for a
# very large file the memory cost isn't worth it for what are, either way,
# best-effort accuracy improvements rather than core extraction.
MAX_BYTES_FOR_FORMULA_AND_MERGE_PASS = 20 * 1024 * 1024


class ExcelExtractionError(Exception):
    """Raised when a workbook cannot be opened or contains no sheets/data."""


@dataclass
class SheetInfo:
    hidden: bool = False
    # A formula cell with no cached value reads as blank in `sheets` (Excel
    # never calculated it, or it was written by a tool that skips caching) —
    # this count means that blank is *not* necessarily a genuinely empty
    # cell, so callers can warn instead of trusting it silently.
    formula_without_cache_count: int = 0


@dataclass
class ExcelExtractionResult:
    sheets: Dict[str, SheetRows] = field(default_factory=dict)
    sheet_info: Dict[str, SheetInfo] = field(default_factory=dict)


def _extract_xlsx(file_path: str) -> ExcelExtractionResult:
    try:
        workbook = openpyxl.load_workbook(file_path, data_only=True, read_only=True)
    except Exception as exc:  # noqa: BLE001
        raise ExcelExtractionError(f'Unable to read the Excel file: {exc}') from exc

    if not workbook.sheetnames:
        raise ExcelExtractionError('The workbook contains no worksheets.')

    sheets: Dict[str, SheetRows] = {}
    sheet_info: Dict[str, SheetInfo] = {}
    for sheet_name in workbook.sheetnames:
        worksheet = workbook[sheet_name]
        rows = [list(row) for row in worksheet.iter_rows(values_only=True)]
        while rows and is_blank_row(rows[-1]):
            rows.pop()
        sheets[sheet_name] = rows
        # sheet_state ('visible'/'hidden'/'veryHidden') is available even in
        # read-only mode — a hidden sheet is still fully processed below
        # (never silently skipped), just flagged so a reviewer knows it
        # wasn't meant to be seen in the original workbook.
        sheet_info[sheet_name] = SheetInfo(hidden=worksheet.sheet_state != 'visible')
    workbook.close()

    if os.path.getsize(file_path) <= MAX_BYTES_FOR_FORMULA_AND_MERGE_PASS:
        _fill_merged_cells_and_flag_formulas(file_path, sheets, sheet_info)

    return ExcelExtractionResult(sheets=sheets, sheet_info=sheet_info)


def _fill_merged_cells_and_flag_formulas(
    file_path: str, sheets: Dict[str, SheetRows], sheet_info: Dict[str, SheetInfo],
) -> None:
    """Best-effort only: merged-cell ranges and raw formula strings are both
    unavailable in the read-only mode used above (openpyxl's
    ``ReadOnlyWorksheet`` has no ``merged_cells`` at all, and a
    ``data_only=True`` load discards the formula source entirely) — this
    does one additional non-streaming, ``data_only=False`` load purely to
    recover those two things. Never raises: a failure here just means these
    two accuracy improvements are skipped, not that extraction fails.
    """
    try:
        workbook = openpyxl.load_workbook(file_path, data_only=False, read_only=False)
    except Exception:  # noqa: BLE001
        return

    try:
        for sheet_name, rows in sheets.items():
            if sheet_name not in workbook.sheetnames or not rows:
                continue
            worksheet = workbook[sheet_name]
            info = sheet_info.setdefault(sheet_name, SheetInfo())

            formula_missing = 0
            for row in worksheet.iter_rows():
                for cell in row:
                    if not (isinstance(cell.value, str) and cell.value.startswith('=')):
                        continue
                    r, c = cell.row - 1, cell.column - 1
                    if r < len(rows) and c < len(rows[r]) and rows[r][c] is None:
                        formula_missing += 1
            info.formula_without_cache_count = formula_missing

            # Only the merged range's top-left ("anchor") cell holds a
            # value in the underlying file — every other cell in the range
            # already came back as None from the values_only read above, so
            # here that value is copied across the rest of the range,
            # exactly matching what Excel visually shows the user.
            for merged_range in worksheet.merged_cells.ranges:
                min_row, min_col = merged_range.min_row, merged_range.min_col
                max_row, max_col = merged_range.max_row, merged_range.max_col
                if min_row - 1 >= len(rows) or min_col - 1 >= len(rows[min_row - 1]):
                    continue
                anchor_value = rows[min_row - 1][min_col - 1]
                if anchor_value is None:
                    continue
                for r in range(min_row, max_row + 1):
                    if r - 1 >= len(rows):
                        continue
                    row_values = rows[r - 1]
                    for c in range(min_col, max_col + 1):
                        if c - 1 >= len(row_values):
                            continue
                        if row_values[c - 1] is None:
                            row_values[c - 1] = anchor_value
    finally:
        workbook.close()


def _xls_cell_value(cell, workbook):
    """xlrd hands back a raw ``Cell`` with a numeric type code rather than a
    native Python value — dates and booleans need decoding, everything else
    (text/number/blank) is already the right Python type."""
    import xlrd

    if cell.ctype == xlrd.XL_CELL_EMPTY:
        return None
    if cell.ctype == xlrd.XL_CELL_DATE:
        try:
            from datetime import datetime
            return datetime(*xlrd.xldate_as_tuple(cell.value, workbook.datemode))
        except (ValueError, OverflowError):
            return cell.value
    if cell.ctype == xlrd.XL_CELL_BOOLEAN:
        return bool(cell.value)
    return cell.value


def _extract_xls(file_path: str) -> ExcelExtractionResult:
    """Legacy binary .xls workbooks predate openpyxl's format and need the
    separate xlrd engine (openpyxl only ever reads the newer XML-based
    .xlsx/.xlsm formats and silently fails, or errors, on a true .xls).
    xlrd exposes neither cached-formula detection nor merged-cell ranges in
    a form worth the added complexity for a legacy format — this is a known,
    documented limitation, not a silent gap (see excel_extractor module
    tests for exactly what .xls support does and doesn't cover)."""
    try:
        import xlrd
    except ImportError as exc:
        raise ExcelExtractionError(
            'Reading legacy .xls files requires the "xlrd" package, which is not installed on this server.'
        ) from exc

    try:
        workbook = xlrd.open_workbook(file_path)
    except Exception as exc:  # noqa: BLE001
        raise ExcelExtractionError(f'Unable to read the Excel file: {exc}') from exc

    if not workbook.sheet_names():
        raise ExcelExtractionError('The workbook contains no worksheets.')

    sheets: Dict[str, SheetRows] = {}
    sheet_info: Dict[str, SheetInfo] = {}
    for sheet_name in workbook.sheet_names():
        sheet = workbook.sheet_by_name(sheet_name)
        rows = [
            [_xls_cell_value(sheet.cell(row_idx, col_idx), workbook) for col_idx in range(sheet.ncols)]
            for row_idx in range(sheet.nrows)
        ]
        while rows and is_blank_row(rows[-1]):
            rows.pop()
        sheets[sheet_name] = rows
        sheet_info[sheet_name] = SheetInfo(hidden=sheet.visibility != 0)
    return ExcelExtractionResult(sheets=sheets, sheet_info=sheet_info)


def extract_excel(file_path: str) -> ExcelExtractionResult:
    ext = os.path.splitext(file_path)[1].lower()
    if ext != '.xls':
        check_zip_bomb(file_path, ExcelExtractionError)
    result = _extract_xls(file_path) if ext == '.xls' else _extract_xlsx(file_path)

    if not any(result.sheets.values()):
        raise ExcelExtractionError('The workbook does not contain any data.')

    return result
